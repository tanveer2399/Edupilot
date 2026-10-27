import csv
import io
import logging
import threading
import time
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from uuid import UUID, uuid4
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from app.models.academic_structure import (
    AcademicYear,
    Branch,
    ClassSection,
    Department,
    FacultySectionAssignment,
    Semester,
    StudentEnrollment,
    Subject,
)
from app.models.student import Student
from app.schemas.faculty import (
    AcademicYearCreate,
    BranchCreate,
    ClassSectionCreate,
    DepartmentCreate,
    SemesterCreate,
    StudentCreate,
    SubjectCreate,
)

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
MAX_IMPORT_COLUMNS = 100
MAX_PREVIEW_RECORDS = 100
PREVIEW_TTL_SECONDS = 15 * 60
MAX_ACTIVE_PREVIEWS = 20
logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class ImportSpec:
    fields: tuple[str, ...]
    required_fields: tuple[str, ...]


IMPORT_SPECS: dict[str, ImportSpec] = {
    "students": ImportSpec(
        (
            "student_id",
            "full_name",
            "email",
            "branch_code",
            "academic_year",
            "semester_number",
            "section_name",
        ),
        (
            "student_id",
            "full_name",
            "email",
            "branch_code",
            "academic_year",
            "semester_number",
            "section_name",
        ),
    ),
    "academic_years": ImportSpec(
        ("name", "start_date", "end_date"),
        ("name", "start_date", "end_date"),
    ),
    "departments": ImportSpec(("code", "name"), ("code", "name")),
    "branches": ImportSpec(
        ("department_code", "code", "name"),
        ("department_code", "code", "name"),
    ),
    "semesters": ImportSpec(
        ("academic_year", "number", "name", "start_date", "end_date"),
        ("academic_year", "number", "name"),
    ),
    "class_sections": ImportSpec(
        ("branch_code", "academic_year", "semester_number", "name"),
        ("branch_code", "academic_year", "semester_number", "name"),
    ),
    "subjects": ImportSpec(
        (
            "branch_code",
            "academic_year",
            "semester_number",
            "code",
            "name",
            "credits",
        ),
        (
            "branch_code",
            "academic_year",
            "semester_number",
            "code",
            "name",
            "credits",
        ),
    ),
    "enrollments": ImportSpec(
        (
            "student_id",
            "branch_code",
            "academic_year",
            "semester_number",
            "section_name",
            "status",
        ),
        (
            "student_id",
            "branch_code",
            "academic_year",
            "semester_number",
            "section_name",
        ),
    ),
}

_previews: dict[UUID, dict[str, Any]] = {}
_previews_lock = threading.Lock()


class ImportFileError(ValueError):
    pass


class ImportCommitError(RuntimeError):
    pass


def import_fields(import_type: str) -> dict[str, list[str]]:
    spec = IMPORT_SPECS.get(import_type)
    if spec is None:
        raise ImportFileError("Unsupported import type.")
    return {
        "fields": list(spec.fields),
        "required_fields": list(spec.required_fields),
    }


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float, bool, Decimal)):
        return value
    return str(value).strip()


def _read_csv(content: bytes) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    try:
        text = content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise ImportFileError("CSV files must use UTF-8 encoding.") from error
    if "\x00" in text:
        raise ImportFileError("CSV file contains unsupported null bytes.")
    try:
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        raw_headers = next(reader, None)
        if raw_headers is None:
            raise ImportFileError("The file is empty.")
        headers = [header.strip() for header in raw_headers]
        rows = []
        for line_number, row in enumerate(reader, start=2):
            if not any(value.strip() for value in row):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ImportFileError(
                    f"Files may contain at most {MAX_IMPORT_ROWS} data rows."
                )
            rows.append((line_number, [_clean_cell(value) for value in row]))
    except csv.Error as error:
        raise ImportFileError("CSV file has invalid quoting or delimiters.") from error
    return headers, rows


def _read_xlsx(content: bytes) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ImportFileError("The Excel file is invalid or corrupt.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 1000
                or sum(member.file_size for member in members)
                > 25 * 1024 * 1024
            ):
                raise ImportFileError(
                    "The Excel workbook expands beyond the allowed size."
                )
    except zipfile.BadZipFile as error:
        raise ImportFileError("The Excel file is invalid or corrupt.") from error
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (
        OSError,
        ValueError,
        KeyError,
        IndexError,
        EOFError,
        zipfile.BadZipFile,
        InvalidFileException,
        ParseError,
    ) as error:
        raise ImportFileError("The Excel file is invalid or cannot be read.") from error

    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ImportFileError("The Excel workbook has no active worksheet.")
        rows = worksheet.iter_rows(values_only=False)
        first_row = next(rows, None)
        if first_row is None:
            raise ImportFileError("The Excel worksheet is empty.")
        headers = []
        if len(first_row) > MAX_IMPORT_COLUMNS:
            raise ImportFileError(
                f"Files may contain at most {MAX_IMPORT_COLUMNS} columns."
            )
        for cell in first_row:
            if cell.data_type == "f":
                raise ImportFileError("Formula cells are not allowed in the header.")
            value = _clean_cell(cell.value)
            headers.append("" if value is None else str(value))
        parsed_rows: list[tuple[int, list[Any]]] = []
        for row_number, cells in enumerate(rows, start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                raise ImportFileError(
                    f"Files may contain at most {MAX_IMPORT_ROWS} data rows."
                )
            if any(cell.data_type == "f" for cell in cells):
                raise ImportFileError(
                    f"Formula cells are not allowed (worksheet row {row_number})."
                )
            values = [_clean_cell(cell.value) for cell in cells]
            if any(value not in (None, "") for value in values):
                parsed_rows.append((row_number, values))
        return headers, parsed_rows
    finally:
        workbook.close()


def read_import_file(
    filename: str | None,
    content: bytes,
) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    if not content:
        raise ImportFileError("Choose a non-empty CSV or Excel file.")
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportFileError("File exceeds the 5 MB upload limit.")
    extension = (filename or "").rsplit(".", 1)[-1].lower()
    if extension == "csv":
        headers, rows = _read_csv(content)
    elif extension == "xlsx":
        headers, rows = _read_xlsx(content)
    else:
        raise ImportFileError("Only .csv and .xlsx files are supported.")
    if not headers or all(not header for header in headers):
        raise ImportFileError("The first row must contain column headers.")
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise ImportFileError(
            f"Files may contain at most {MAX_IMPORT_COLUMNS} columns."
        )
    normalized = [header.casefold() for header in headers]
    if len(set(normalized)) != len(normalized):
        raise ImportFileError("Column headers must be unique (ignoring case).")
    if any(not header for header in headers):
        raise ImportFileError("Column headers cannot be blank.")
    if len(rows) > MAX_IMPORT_ROWS:
        raise ImportFileError(
            f"Files may contain at most {MAX_IMPORT_ROWS} data rows."
        )
    if any(len(row) > len(headers) for _, row in rows):
        raise ImportFileError("A data row contains more cells than the header.")
    return headers, rows


def _preview_samples(
    headers: list[str],
    rows: list[tuple[int, list[Any]]],
) -> list[dict[str, Any]]:
    return [
        {
            "row_number": row_number,
            "values": {
                headers[index]: value
                for index, value in enumerate(values[: len(headers)])
            },
        }
        for row_number, values in rows[:5]
    ]


def _expire_previews(now: float) -> None:
    expired = [
        preview_id
        for preview_id, preview in _previews.items()
        if now - preview["created_at"] > PREVIEW_TTL_SECONDS
    ]
    for preview_id in expired:
        del _previews[preview_id]
    while len(_previews) > MAX_ACTIVE_PREVIEWS:
        _previews.pop(next(iter(_previews)))


def _as_text(row: dict[str, Any], field: str) -> str:
    value = row.get(field)
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field} is required.")
    text = str(value).strip()
    if len(text) > 500:
        raise ValueError(f"{field} exceeds the 500 character import limit.")
    return text


def _as_int(row: dict[str, Any], field: str) -> int:
    value = row.get(field)
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, AttributeError) as error:
        raise ValueError(f"{field} must be a whole number.") from error
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f"{field} must be a whole number.")
    return int(number)


def _optional_date(row: dict[str, Any], field: str) -> date | None:
    value = row.get(field)
    if value is None or str(value).strip() == "":
        return None
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError as error:
        raise ValueError(f"{field} must be a valid YYYY-MM-DD date.") from error


def _required_date(row: dict[str, Any], field: str) -> date:
    value = _optional_date(row, field)
    if value is None:
        raise ValueError(f"{field} is required.")
    return value


def _mapped_rows(
    headers: list[str],
    rows: list[tuple[int, list[Any]]],
    import_type: str,
    mapping: dict[str, str],
) -> list[tuple[int, dict[str, Any]]]:
    spec = IMPORT_SPECS[import_type]
    missing_fields = sorted(set(spec.required_fields) - set(mapping))
    invalid_fields = sorted(set(mapping) - set(spec.fields))
    missing_headers = sorted(set(mapping.values()) - set(headers))
    if missing_fields:
        raise ImportFileError(
            f"Map required fields: {', '.join(missing_fields)}."
        )
    if invalid_fields:
        raise ImportFileError(f"Unknown mapped fields: {', '.join(invalid_fields)}.")
    if missing_headers:
        raise ImportFileError(
            f"Mapped source columns are missing: {', '.join(missing_headers)}."
        )
    mapped_headers = list(mapping.values())
    if len(set(mapped_headers)) != len(mapped_headers):
        raise ImportFileError("Each source column can map to only one field.")
    header_indexes = {header: index for index, header in enumerate(headers)}
    mapped: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in rows:
        record = {}
        for field, header in mapping.items():
            index = header_indexes[header]
            record[field] = values[index] if index < len(values) else None
        mapped.append((row_number, record))
    return mapped


def _owned_department(db: Session, code: str, user_id: str) -> Department:
    record = db.scalar(
        select(Department).where(
            Department.code == code.upper(),
            Department.created_by_user_id == user_id,
        )
    )
    if record is None:
        raise ValueError(f"Department code '{code}' is not in your faculty scope.")
    return record


def _owned_year(db: Session, name: str, user_id: str) -> AcademicYear:
    record = db.scalar(
        select(AcademicYear).where(
            AcademicYear.name == name,
            AcademicYear.created_by_user_id == user_id,
        )
    )
    if record is None:
        raise ValueError(f"Academic year '{name}' is not in your faculty scope.")
    return record


def _owned_branch(db: Session, code: str, user_id: str) -> Branch:
    records = list(
        db.scalars(
            select(Branch)
            .where(
                Branch.code == code.upper(),
                Branch.created_by_user_id == user_id,
            )
            .order_by(Branch.id)
        )
    )
    if not records:
        raise ValueError(f"Branch code '{code}' is not in your faculty scope.")
    if len(records) > 1:
        raise ValueError(
            f"Branch code '{code}' is ambiguous; branch codes must be unique "
            "within your faculty scope for imports."
        )
    return records[0]


def _owned_semester(
    db: Session,
    year_id: int,
    number: int,
    user_id: str,
) -> Semester:
    record = db.scalar(
        select(Semester).where(
            Semester.academic_year_id == year_id,
            Semester.number == number,
            Semester.created_by_user_id == user_id,
        )
    )
    if record is None:
        raise ValueError(
            f"Semester {number} for the selected academic year is not in your faculty scope."
        )
    return record


def _owned_section(
    db: Session,
    branch_id: int,
    year_id: int,
    semester_id: int,
    name: str,
    user_id: str,
) -> ClassSection:
    record = db.scalar(
        select(ClassSection)
        .join(
            FacultySectionAssignment,
            FacultySectionAssignment.section_id == ClassSection.id,
        )
        .where(
            ClassSection.branch_id == branch_id,
            ClassSection.academic_year_id == year_id,
            ClassSection.semester_id == semester_id,
            ClassSection.name == name,
            FacultySectionAssignment.faculty_user_id == user_id,
        )
    )
    if record is None:
        raise ValueError(
            f"Class section '{name}' is not assigned to your faculty account."
        )
    return record


def _duplicate_lookup(
    db: Session,
    import_type: str,
    values: dict[str, Any],
    user_id: str,
) -> tuple[Any | None, str]:
    model = None
    statement = None
    identity = ""
    if import_type == "academic_years":
        model = AcademicYear
        statement = select(model).where(model.name == values["name"])
        identity = values["name"]
    elif import_type == "departments":
        model = Department
        statement = select(model).where(Department.code == values["code"])
        identity = values["code"]
    elif import_type == "branches":
        model = Branch
        statement = select(model).where(
            Branch.department_id == values["department_id"],
            Branch.code == values["code"],
        )
        identity = f"{values['department_id']}:{values['code']}"
    elif import_type == "semesters":
        model = Semester
        statement = select(model).where(
            Semester.academic_year_id == values["academic_year_id"],
            Semester.number == values["number"],
        )
        identity = f"{values['academic_year_id']}:{values['number']}"
    elif import_type == "class_sections":
        model = ClassSection
        statement = select(model).where(
            ClassSection.branch_id == values["branch_id"],
            ClassSection.academic_year_id == values["academic_year_id"],
            ClassSection.semester_id == values["semester_id"],
            ClassSection.name == values["name"],
        )
        identity = (
            f"{values['branch_id']}:{values['academic_year_id']}:"
            f"{values['semester_id']}:{values['name']}"
        )
    elif import_type == "subjects":
        model = Subject
        statement = select(model).where(
            Subject.branch_id == values["branch_id"],
            Subject.semester_id == values["semester_id"],
            Subject.code == values["code"],
        )
        identity = f"{values['branch_id']}:{values['semester_id']}:{values['code']}"
    elif import_type == "students":
        model = Student
        statement = select(model).where(Student.student_id == values["student_id"])
        identity = values["student_id"]
    elif import_type == "enrollments":
        model = StudentEnrollment
        statement = select(model).where(
            StudentEnrollment.student_id == values["student_id"],
            StudentEnrollment.semester_id == values["semester_id"],
        )
        identity = f"{values['student_id']}:{values['semester_id']}"
    if statement is None or model is None:
        raise ImportFileError("Unsupported import type.")
    existing = db.scalar(statement)
    if existing is not None:
        if import_type == "enrollments":
            authorized = db.scalar(
                select(FacultySectionAssignment.section_id).where(
                    FacultySectionAssignment.section_id == existing.section_id,
                    FacultySectionAssignment.faculty_user_id == user_id,
                )
            )
            if authorized is None:
                raise ValueError("The existing enrollment is outside your faculty scope.")
        elif import_type == "students":
            owned = db.scalar(
                select(StudentEnrollment.id)
                .join(
                    FacultySectionAssignment,
                    FacultySectionAssignment.section_id == StudentEnrollment.section_id,
                )
                .where(
                    StudentEnrollment.student_id == existing.student_id,
                    FacultySectionAssignment.faculty_user_id == user_id,
                )
            )
            if owned is None:
                raise ValueError("The existing student is outside your faculty scope.")
        elif getattr(existing, "created_by_user_id", None) != user_id:
            raise ValueError("The matching academic record is outside your faculty scope.")
    return existing, identity


def _build_values(
    db: Session,
    import_type: str,
    row: dict[str, Any],
    user_id: str,
) -> tuple[dict[str, Any], Any | None, str]:
    if import_type == "academic_years":
        payload = {
            "name": _as_text(row, "name"),
            "start_date": _required_date(row, "start_date"),
            "end_date": _required_date(row, "end_date"),
        }
        checked = AcademicYearCreate.model_validate(payload)
        values = checked.model_dump()
    elif import_type == "departments":
        checked = DepartmentCreate.model_validate(
            {"code": _as_text(row, "code"), "name": _as_text(row, "name")}
        )
        values = checked.model_dump()
        values["code"] = checked.code.upper()
    elif import_type == "branches":
        department = _owned_department(db, _as_text(row, "department_code"), user_id)
        checked = BranchCreate.model_validate(
            {
                "department_id": department.id,
                "code": _as_text(row, "code"),
                "name": _as_text(row, "name"),
            }
        )
        values = checked.model_dump()
        values["code"] = checked.code.upper()
    elif import_type == "semesters":
        year = _owned_year(db, _as_text(row, "academic_year"), user_id)
        checked = SemesterCreate.model_validate(
            {
                "academic_year_id": year.id,
                "number": _as_int(row, "number"),
                "name": _as_text(row, "name"),
                "start_date": _optional_date(row, "start_date"),
                "end_date": _optional_date(row, "end_date"),
            }
        )
        values = checked.model_dump()
    elif import_type == "class_sections":
        branch = _owned_branch(db, _as_text(row, "branch_code"), user_id)
        year = _owned_year(db, _as_text(row, "academic_year"), user_id)
        semester = _owned_semester(
            db,
            year.id,
            _as_int(row, "semester_number"),
            user_id,
        )
        checked = ClassSectionCreate.model_validate(
            {
                "branch_id": branch.id,
                "academic_year_id": year.id,
                "semester_id": semester.id,
                "name": _as_text(row, "name"),
            }
        )
        values = checked.model_dump()
    elif import_type == "subjects":
        branch = _owned_branch(db, _as_text(row, "branch_code"), user_id)
        year = _owned_year(db, _as_text(row, "academic_year"), user_id)
        semester = _owned_semester(
            db,
            year.id,
            _as_int(row, "semester_number"),
            user_id,
        )
        try:
            credits = Decimal(str(row.get("credits")).strip())
        except (InvalidOperation, AttributeError) as error:
            raise ValueError("credits must be a positive decimal number.") from error
        checked = SubjectCreate.model_validate(
            {
                "branch_id": branch.id,
                "semester_id": semester.id,
                "code": _as_text(row, "code"),
                "name": _as_text(row, "name"),
                "credits": credits,
            }
        )
        values = checked.model_dump()
        values["code"] = checked.code.upper()
    elif import_type == "students":
        section = _section_from_row(db, row, user_id)
        checked = StudentCreate.model_validate(
            {
                "student_id": _as_text(row, "student_id"),
                "full_name": _as_text(row, "full_name"),
                "email": _as_text(row, "email"),
                "section_id": section.id,
            }
        )
        values = checked.model_dump()
        values["email"] = str(checked.email).casefold()
        branch = db.get(Branch, section.branch_id)
        semester = db.get(Semester, section.semester_id)
        if branch is None:
            raise ValueError("The selected class section has no valid branch.")
        if semester is None:
            raise ValueError("The selected class section has no valid semester.")
        values["department"] = branch.name
        values["semester"] = semester.number
        email_match = db.scalar(
            select(Student.student_id).where(
                Student.email == values["email"],
                Student.student_id != values["student_id"],
            )
        )
        if email_match is not None:
            raise ValueError(
                f"Email '{values['email']}' is already used by another student."
            )
    else:
        student_id = _as_text(row, "student_id")
        student = db.get(Student, student_id)
        if student is None:
            raise ValueError(
                f"Student ID '{student_id}' does not exist; import the student first."
            )
        section = _section_from_row(db, row, user_id)
        status_value = str(row.get("status") or "active").strip().lower()
        if status_value not in {"active", "completed", "withdrawn"}:
            raise ValueError("status must be active, completed, or withdrawn.")
        values = {
            "student_id": student_id,
            "section_id": section.id,
            "semester_id": section.semester_id,
            "status": status_value,
        }

    existing, identity = _duplicate_lookup(db, import_type, values, user_id)
    return values, existing, identity





def _section_from_row(
    db: Session,
    row: dict[str, Any],
    user_id: str,
) -> ClassSection:
    branch = _owned_branch(db, _as_text(row, "branch_code"), user_id)
    year = _owned_year(db, _as_text(row, "academic_year"), user_id)
    semester = _owned_semester(
        db,
        year.id,
        _as_int(row, "semester_number"),
        user_id,
    )
    return _owned_section(
        db,
        branch.id,
        year.id,
        semester.id,
        _as_text(row, "section_name"),
        user_id,
    )


def _validate_rows(
    db: Session,
    import_type: str,
    rows: list[tuple[int, dict[str, Any]]],
    user_id: str,
    duplicate_policy: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    planned: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    for row_number, row in rows:
        try:
            values, existing, identity = _build_values(
                db,
                import_type,
                row,
                user_id,
            )
            prior_row = seen.get(identity)
            if prior_row is not None:
                raise ValueError(
                    f"Duplicate identifier also appears in file row {prior_row}."
                )
            seen[identity] = row_number
            if existing is None:
                action = "created"
            elif duplicate_policy == "update":
                action = "updated"
            else:
                action = "skipped"
            preview_values = {
                key: _clean_cell(value)
                for key, value in values.items()
                if key not in {"department", "semester"}
            }
            planned.append(
                {
                    "row_number": row_number,
                    "identity": identity,
                    "action": action,
                    "values": values,
                    "existing_id": (
                        existing.id
                        if existing is not None and hasattr(existing, "id")
                        else None
                    ),
                    "expected_state": (
                        _current_state(
                            db,
                            import_type,
                            existing,
                            values,
                        )
                        if existing is not None
                        else None
                    ),
                    "before": (
                        _public_existing_values(import_type, existing)
                        if existing is not None
                        else None
                    ),
                    "preview": preview_values,
                }
            )
        except (ValueError, ValidationError) as error:
            if isinstance(error, ValidationError):
                detail = "; ".join(
                    f"{'.'.join(str(part) for part in item['loc'])}: {item['msg']}"
                    for item in error.errors(include_input=False)
                )
            else:
                detail = str(error)
            errors.append({"row_number": row_number, "message": detail})
    return planned, errors


def _stable_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _current_state(
    db: Session,
    import_type: str,
    record: Any,
    values: dict[str, Any],
) -> dict[str, Any]:
    state = {
        column.name: _stable_value(getattr(record, column.name))
        for column in record.__table__.columns
    }
    if import_type == "students":
        section = db.get(ClassSection, values["section_id"])
        if section is None:
            raise ValueError("The selected class section no longer exists.")
        enrollment = db.scalar(
            select(StudentEnrollment).where(
                StudentEnrollment.student_id == record.student_id,
                StudentEnrollment.semester_id == section.semester_id,
            )
        )
        state["enrollment"] = (
            {
                column.name: _stable_value(getattr(enrollment, column.name))
                for column in enrollment.__table__.columns
            }
            if enrollment is not None
            else None
        )
    return state


def _public_existing_values(import_type: str, record: Any) -> dict[str, Any]:
    fields_by_type = {
        "academic_years": ("name", "start_date", "end_date"),
        "departments": ("code", "name"),
        "branches": ("department_id", "code", "name"),
        "semesters": ("academic_year_id", "number", "name", "start_date", "end_date"),
        "class_sections": ("branch_id", "academic_year_id", "semester_id", "name"),
        "subjects": ("branch_id", "semester_id", "code", "name", "credits"),
        "students": ("student_id", "full_name", "email", "department", "semester"),
        "enrollments": ("student_id", "section_id", "semester_id", "status"),
    }
    return {
        field: _clean_cell(getattr(record, field))
        for field in fields_by_type[import_type]
    }


def prepare_import_preview(
    db: Session,
    *,
    filename: str | None,
    content: bytes,
    import_type: str,
    mapping: dict[str, str] | None,
    duplicate_policy: str,
    user_id: str,
) -> dict[str, Any]:
    if import_type not in IMPORT_SPECS:
        raise ImportFileError("Unsupported import type.")
    if duplicate_policy not in {"skip", "update"}:
        raise ImportFileError("duplicate_policy must be 'skip' or 'update'.")
    headers, file_rows = read_import_file(filename, content)
    base_response = {
        "import_type": import_type,
        "headers": headers,
        "mapping_fields": import_fields(import_type),
        "sample_rows": _preview_samples(headers, file_rows),
        "row_count": len(file_rows),
    }
    if mapping is None:
        return {**base_response, "requires_mapping": True}
    mapped_rows = _mapped_rows(headers, file_rows, import_type, mapping)
    planned, errors = _validate_rows(
        db,
        import_type,
        mapped_rows,
        user_id,
        duplicate_policy,
    )
    preview_id = uuid4()
    summary = {
        "created": sum(item["action"] == "created" for item in planned),
        "updated": sum(item["action"] == "updated" for item in planned),
        "skipped": sum(item["action"] == "skipped" for item in planned),
        "failed": len(errors),
    }
    response_records = [
        {
            "row_number": item["row_number"],
            "identity": item["identity"],
            "action": item["action"],
            "values": item["preview"],
            "before": item["before"],
        }
        for item in planned[:MAX_PREVIEW_RECORDS]
    ]
    response = {
        **base_response,
        "requires_mapping": False,
        "preview_id": str(preview_id),
        "duplicate_policy": duplicate_policy,
        "summary": summary,
        "errors": errors[:MAX_PREVIEW_RECORDS],
        "errors_truncated": len(errors) > MAX_PREVIEW_RECORDS,
        "records": response_records,
        "records_truncated": len(planned) > MAX_PREVIEW_RECORDS,
        "can_commit": bool(planned) and not errors,
    }
    if not errors and planned:
        with _previews_lock:
            _expire_previews(time.monotonic())
            _previews[preview_id] = {
                "created_at": time.monotonic(),
                "user_id": user_id,
                "import_type": import_type,
                "duplicate_policy": duplicate_policy,
                "planned": planned,
            }
    return response


def _record_for_existing(
    db: Session,
    import_type: str,
    operation: dict[str, Any],
    user_id: str,
) -> Any:
    values = operation["values"]
    if import_type == "academic_years":
        return _owned_record(db, AcademicYear, operation["existing_id"], user_id)
    if import_type == "departments":
        return _owned_record(db, Department, operation["existing_id"], user_id)
    if import_type == "branches":
        return _owned_record(db, Branch, operation["existing_id"], user_id)
    if import_type == "semesters":
        return _owned_record(db, Semester, operation["existing_id"], user_id)
    if import_type == "class_sections":
        assignment = db.scalar(
            select(FacultySectionAssignment.section_id).where(
                FacultySectionAssignment.section_id == operation["existing_id"],
                FacultySectionAssignment.faculty_user_id == user_id,
            )
        )
        record = db.get(ClassSection, operation["existing_id"])
        if assignment is None or record is None:
            raise ImportCommitError("The class section is no longer in your scope.")
        return record
    if import_type == "subjects":
        return _owned_record(db, Subject, operation["existing_id"], user_id)
    if import_type == "students":
        record = db.get(Student, values["student_id"])
        if record is None:
            raise ImportCommitError("A student selected for update no longer exists.")
        _section_from_values(db, values, user_id)
        return record
    if import_type == "enrollments":
        record = db.get(StudentEnrollment, operation["existing_id"])
        if record is None:
            raise ImportCommitError("An enrollment selected for update no longer exists.")
        if db.scalar(
            select(FacultySectionAssignment.section_id).where(
                FacultySectionAssignment.section_id == record.section_id,
                FacultySectionAssignment.faculty_user_id == user_id,
            )
        ) is None:
            raise ImportCommitError("The enrollment is no longer in your scope.")
        _section_from_values(db, values, user_id)
        return record
    raise ImportCommitError("Unsupported import type.")


def _owned_record(db: Session, model: type, record_id: int, user_id: str):
    record = db.scalar(
        select(model).where(
            model.id == record_id,
            model.created_by_user_id == user_id,
        )
    )
    if record is None:
        raise ImportCommitError("An academic record is no longer in your scope.")
    return record


def _section_from_values(
    db: Session,
    values: dict[str, Any],
    user_id: str,
) -> ClassSection:
    section = db.get(ClassSection, values["section_id"])
    if section is None or db.scalar(
        select(FacultySectionAssignment.section_id).where(
            FacultySectionAssignment.section_id == section.id,
            FacultySectionAssignment.faculty_user_id == user_id,
        )
    ) is None:
        raise ImportCommitError("The class section is no longer in your scope.")
    return section


def _new_record(
    db: Session,
    import_type: str,
    values: dict[str, Any],
    user_id: str,
) -> Any:
    if import_type == "academic_years":
        return AcademicYear(**values, created_by_user_id=user_id)
    if import_type == "departments":
        return Department(**values, created_by_user_id=user_id)
    if import_type == "branches":
        return Branch(**values, created_by_user_id=user_id)
    if import_type == "semesters":
        return Semester(**values, created_by_user_id=user_id)
    if import_type == "class_sections":
        section = ClassSection(**values, created_by_user_id=user_id)
        db.add(section)
        db.flush()
        db.add(
            FacultySectionAssignment(
                faculty_user_id=user_id,
                section_id=section.id,
            )
        )
        return section
    if import_type == "subjects":
        return Subject(**values, created_by_user_id=user_id)
    if import_type == "students":
        student_values = dict(values)
        section_id = student_values.pop("section_id")
        department = student_values.pop("department")
        semester = student_values.pop("semester")
        student = Student(
            **student_values,
            department=department,
            semester=semester,
            is_demo=False,
        )
        db.add(student)
        db.flush()
        section = _section_from_values(db, {"section_id": section_id}, user_id)
        db.add(
            StudentEnrollment(
                student_id=student.student_id,
                section_id=section.id,
                semester_id=section.semester_id,
            )
        )
        return student
    if import_type == "enrollments":
        return StudentEnrollment(**values)
    raise ImportCommitError("Unsupported import type.")


def _update_record(
    db: Session,
    import_type: str,
    existing: Any,
    values: dict[str, Any],
    user_id: str,
) -> None:
    if import_type in {"academic_years", "departments", "branches", "semesters", "subjects"}:
        for field, value in values.items():
            setattr(existing, field, value)
        return
    if import_type == "students":
        existing.full_name = values["full_name"]
        existing.email = values["email"]
        existing.department = values["department"]
        existing.semester = values["semester"]
        section = _section_from_values(
            db,
            {"section_id": values["section_id"]},
            user_id,
        )
        enrollment = db.scalar(
            select(StudentEnrollment)
            .join(
                FacultySectionAssignment,
                FacultySectionAssignment.section_id == StudentEnrollment.section_id,
            )
            .where(
                StudentEnrollment.student_id == existing.student_id,
                StudentEnrollment.semester_id == section.semester_id,
                FacultySectionAssignment.faculty_user_id == user_id,
            )
        )
        if enrollment is None:
            enrollment = StudentEnrollment(
                student_id=existing.student_id,
                section_id=section.id,
                semester_id=section.semester_id,
            )
            db.add(enrollment)
        else:
            enrollment.section_id = section.id
            enrollment.semester_id = section.semester_id
        return
    if import_type == "enrollments":
        existing.section_id = values["section_id"]
        existing.semester_id = values["semester_id"]
        existing.status = values["status"]


def commit_import(
    db: Session,
    *,
    preview_id: UUID,
    user_id: str,
) -> dict[str, int]:
    with _previews_lock:
        _expire_previews(time.monotonic())
        preview = _previews.get(preview_id)
        if preview is None or preview["user_id"] != user_id:
            raise ImportFileError("Import preview not found or expired.")
        del _previews[preview_id]

    try:
        for operation in preview["planned"]:
            if operation["action"] == "skipped":
                continue
            if operation["action"] == "created":
                existing, _ = _duplicate_lookup(
                    db,
                    preview["import_type"],
                    operation["values"],
                    user_id,
                )
                if existing is not None:
                    raise ImportCommitError(
                        "A duplicate was added after preview. Create a new preview."
                    )
                record = _new_record(
                    db,
                    preview["import_type"],
                    operation["values"],
                    user_id,
                )
                db.add(record)
            else:
                existing = _record_for_existing(
                    db,
                    preview["import_type"],
                    operation,
                    user_id,
                )
                current_state = _current_state(
                    db,
                    preview["import_type"],
                    existing,
                    operation["values"],
                )
                if current_state != operation["expected_state"]:
                    raise ImportCommitError(
                        "A record changed after preview. Create a new preview and review it again."
                    )
                _update_record(
                    db,
                    preview["import_type"],
                    existing,
                    operation["values"],
                    user_id,
                )
            db.flush()
        db.commit()
    except (IntegrityError, SQLAlchemyError, ImportCommitError) as error:
        db.rollback()
        if isinstance(error, ImportCommitError):
            raise ImportCommitError(str(error)) from error
        logger.warning("Faculty import transaction failed: %s", type(error).__name__)
        raise ImportCommitError(
            "The import could not be saved. No rows were committed; review the file and try again."
        ) from error
    summary = {
        "created": sum(item["action"] == "created" for item in preview["planned"]),
        "updated": sum(item["action"] == "updated" for item in preview["planned"]),
        "skipped": sum(item["action"] == "skipped" for item in preview["planned"]),
        "failed": 0,
    }
    return summary
